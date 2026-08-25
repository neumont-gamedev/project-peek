"""Generate the Word tracker (PRO100_Team_Tracker.docx) from teams_data.py + history.json.

Run:  python build_doc.py
Writes the .docx to the parent Teams folder (one level up from tracker/).
Requires python-docx:  pip install python-docx
"""
import os
import docx
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from common import load_courses, load_history, team_activity, STATUS_LABEL, STATUS_ORDER

HERE = os.path.dirname(os.path.abspath(__file__))
OUTDIR = os.path.abspath(os.path.join(HERE, ".."))

STATUS_RGB = {
    "healthy": RGBColor(0x2E, 0x7D, 0x32),
    "watch":   RGBColor(0xB5, 0x6A, 0x00),
    "risk":    RGBColor(0xB0, 0x00, 0x00),
    "noaccess": RGBColor(0x80, 0x80, 0x80),
}
RED = RGBColor(0xB0, 0x00, 0x00)


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'), r_id)
    run = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); rPr.append(color)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    run.append(rPr)
    t = OxmlElement('w:t'); t.text = text; run.append(t)
    hyperlink.append(run); paragraph._p.append(hyperlink)


def label_line(doc, label, value=None, url=None):
    p = doc.add_paragraph(); r = p.add_run(label + ": "); r.bold = True
    if url:
        add_hyperlink(p, url, url)
    elif value:
        p.add_run(value)
    return p


def build_course(course):
    TEAMS = course["teams"]
    REPORT_DATE = course.get("report_date", "")
    course_title = course.get("name", "PRO100 – Introductory Software Projects")
    hist = load_history(course["id"])
    acts = {t["id"]: team_activity(t, hist) for t in TEAMS}

    doc = Document()
    style = doc.styles['Normal']; style.font.name = 'Calibri'; style.font.size = Pt(11)

    # ---- Cover ----
    for _ in range(6): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(course_title); r.font.size = Pt(28); r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Team Progress & Sprint Review Tracker"); r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{len(TEAMS)} teams · as of {REPORT_DATE}"); r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- Activity Dashboard ----
    doc.add_heading("Activity Dashboard", level=1)
    p = doc.add_paragraph(); r = p.add_run(f"GitHub commit activity as of {REPORT_DATE}. Sorted by attention needed.")
    r.italic = True
    rows = sorted(TEAMS, key=lambda t: STATUS_ORDER.index(t["status"]))
    table = doc.add_table(rows=1, cols=6); table.style = 'Light Grid Accent 1'
    for i, lbl in enumerate(["Team", "Commits", "Active days", "Last commit", "Status", "Note"]):
        table.rows[0].cells[i].text = ""; run = table.rows[0].cells[i].paragraphs[0].add_run(lbl); run.bold = True
    for t in rows:
        a = acts[t["id"]]; cells = table.add_row().cells
        cells[0].text = t["team"]
        cells[1].text = str(a["total"]) if a["ok"] else "—"
        cells[2].text = str(a["days"]) if a["ok"] else "—"
        cells[3].text = a["last_fmt"] if a["ok"] else "—"
        cells[4].text = ""; sr = cells[4].paragraphs[0].add_run(STATUS_LABEL[t["status"]])
        sr.bold = True; sr.font.color.rgb = STATUS_RGB[t["status"]]
        cells[5].text = t.get("why", "") or ("Balanced / on track" if t["status"] == "healthy" else "")
    p = doc.add_paragraph(); p.add_run("Legend:  ").bold = True
    for st in ["healthy", "watch", "risk", "noaccess"]:
        run = p.add_run(STATUS_LABEL[st] + "   "); run.bold = True; run.font.color.rgb = STATUS_RGB[st]
    p = doc.add_paragraph()
    nr = p.add_run("Note: commit counts are a discussion signal, not a measure of contribution size. "
                   "Some GitHub usernames were matched to students by inference (marked “likely”).")
    nr.italic = True; nr.font.size = Pt(9); nr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- Roster ----
    doc.add_heading("Team Roster", level=1)
    p = doc.add_paragraph(); p.add_run("Quick reference. Each team has a dedicated page following this table.").italic = True
    table = doc.add_table(rows=1, cols=4); table.style = 'Light Grid Accent 1'
    for i, lbl in enumerate(["Team", "Project", "Members", "Links"]):
        table.rows[0].cells[i].text = ""; run = table.rows[0].cells[i].paragraphs[0].add_run(lbl); run.bold = True
    for t in TEAMS:
        cells = table.add_row().cells
        cells[0].text = t["team"]; cells[1].text = t["project"]; cells[2].text = t["members"]
        lp = cells[3].paragraphs[0]; add_hyperlink(lp, t["github"], "GitHub")
        if t.get("trello"): lp.add_run("  |  "); add_hyperlink(lp, t["trello"], "Trello")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- One page per team ----
    for idx, t in enumerate(TEAMS):
        a = acts[t["id"]]
        doc.add_heading(t["team"], level=1)
        label_line(doc, "Project", t["project"])
        label_line(doc, "Members", t["members"])
        label_line(doc, "GitHub", url=t["github"])
        if t.get("trello"): label_line(doc, "Trello", url=t["trello"])
        else: label_line(doc, "Trello", "(not provided)")
        if t.get("notes"):
            p = doc.add_paragraph(); r = p.add_run("Note: "); r.bold = True; r.font.color.rgb = RED
            rn = p.add_run(t["notes"]); rn.font.color.rgb = RED

        # Project summary
        doc.add_paragraph(); doc.add_heading("Project Summary", level=2)
        label_line(doc, "Concept", t["prop"]["concept"])
        label_line(doc, "Key features", t["prop"]["features"])
        label_line(doc, "Tech", t["prop"]["tech"])
        if t["prop"].get("flag"):
            p = doc.add_paragraph(); r = p.add_run("Heads-up: "); r.bold = True; r.font.color.rgb = RED
            rn = p.add_run(t["prop"]["flag"]); rn.font.color.rgb = RED

        # Sprint planning
        doc.add_paragraph(); doc.add_heading("Sprint 1 Planning", level=2)
        s = t.get("sprint")
        if s:
            label_line(doc, "Product Owner", s["po"])
            label_line(doc, "Scrum Master", s["sm"])
            label_line(doc, "Developers", s["devs"])
            label_line(doc, "Sprint Goal", s["goal"])
            if s.get("hours"): label_line(doc, "Estimated Hours", s["hours"])
        else:
            p = doc.add_paragraph(); r = p.add_run("Sprint 1 planning not submitted."); r.italic = True; r.font.color.rgb = RED

        # GitHub activity
        doc.add_paragraph(); doc.add_heading(f"GitHub Activity · as of {REPORT_DATE}", level=2)
        if a["ok"]:
            label_line(doc, "Total commits", str(a["total"]))
            label_line(doc, "Active window", f"{a['span']}  ({a['days']} active days; last commit {a['last_fmt']})")
            contribs = ", ".join(f"{c['n']} {c['c']} ({c['p']}%{', ' + c['note'] if c['note'] else ''})" for c in a["contribs"])
            label_line(doc, "Contributors", contribs)
            p = doc.add_paragraph(); r = p.add_run("Observation: "); r.bold = True; p.add_run(a["obs"])
        else:
            p = doc.add_paragraph(); r = p.add_run("Not reviewable: "); r.bold = True; r.font.color.rgb = RED
            rn = p.add_run(a["obs"]); rn.font.color.rgb = RED

        # Trello board (only if collected)
        tr = a.get("trello")
        if tr:
            doc.add_paragraph(); doc.add_heading(f"Trello Board · as of {tr.get('date','')}", level=2)
            label_line(doc, "Open cards", f"{tr.get('total_cards', 0)} across {len(tr.get('lists', []))} lists")
            for l in tr.get("lists", []):
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{l.get('name','?')}: ").bold = True
                p.add_run(str(l.get("count", 0)))

        # Review notes
        doc.add_paragraph(); doc.add_heading("Sprint 1 Review Notes", level=2)
        for _ in range(3): doc.add_paragraph("")
        doc.add_heading("General Notes", level=2)
        for _ in range(2): doc.add_paragraph("")

        if idx != len(TEAMS) - 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    out = os.path.join(OUTDIR, course["output_basename"] + ".docx")
    doc.save(out)
    print("Saved:", out)


def build():
    for course in load_courses():
        build_course(course)


if __name__ == "__main__":
    build()
